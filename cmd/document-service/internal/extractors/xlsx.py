import logging
from typing import Tuple
import io

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

logger = logging.getLogger(__name__)

def extract_xlsx(file_obj: io.BytesIO) -> Tuple[str, int]:
    """Extract text from an XLSX file.
    
    Returns:
        Tuple of (extracted_text, number_of_sheets)
    """
    if not HAS_OPENPYXL:
        raise Exception("openpyxl not installed")
    
    try:
        wb = load_workbook(file_obj, data_only=True)
        sheets = len(wb.sheetnames)
        text_parts = []
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text_parts.append(f"\n--- Sheet: {sheet_name} ---")
            
            # Extract all cells with values
            for row in sheet.iter_rows(values_only=True):
                row_values = [str(cell) if cell is not None else "" for cell in row]
                # Only include rows that have at least one non-empty cell
                if any(v.strip() for v in row_values):
                    text_parts.append(" | ".join(row_values))
        
        return "\n".join(text_parts), sheets
    except Exception as e:
        logger.error(f"XLSX extraction failed: {e}")
        raise Exception(f"XLSX extraction failed: {e}")